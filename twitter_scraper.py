import os
import re
import time
import signal
import tempfile
import subprocess
from datetime import datetime
from urllib.parse import urlparse

# ==============================================================================
# CONFIGURATION
# ==============================================================================

# The project folder itself (works on any computer and on GitHub runners).
# The old code pointed at one fixed Windows folder on one laptop; that path
# could never exist on GitHub, so the script could not run there at all.
try:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))
except NameError:
    # Running in Jupyter (code pasted into cells): __file__ does not exist
    # there, so fall back to the notebook's working folder.
    BASE_DIR = os.getcwd()

OUTPUT_FOLDER = os.path.join(BASE_DIR, "Excel_Output")
AUTOMATION_PROFILE_DIR = os.path.join(BASE_DIR, "Chrome_Profile")
TEMP_DIR = os.path.join(BASE_DIR, "Temp")
DEBUG_FOLDER = os.path.join(BASE_DIR, "Debug")

PROFILES_TO_SCRAPE = [
    "https://x.com/mppchaudhary?lang=en",
]

START_DATE = datetime.strptime(
    "2026-10-15 00:00:00", "%Y-%m-%d %H:%M:%S"
)

END_DATE = datetime.strptime(
    "2026-10-21 23:59:59", "%Y-%m-%d %H:%M:%S"
)

# ==============================================================================
# RUN INPUTS FROM GITHUB (optional)
# When the workflow is started from the Actions tab, the values typed into
# the input boxes arrive here as environment variables. Anything typed
# there wins over the defaults written above; anything left blank keeps
# the default from this file. When the script runs on a computer (for
# example in Jupyter), these variables are simply empty and the defaults
# above are used.
# ==============================================================================

def _date_from_input(value, end_of_day=False):
    """Read a date typed into a GitHub input box (date or date + time)."""

    value = (value or "").strip()

    if not value:
        return None

    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            parsed = datetime.strptime(value, fmt)
            if end_of_day and fmt == "%Y-%m-%d":
                parsed = parsed.replace(hour=23, minute=59, second=59)
            return parsed
        except ValueError:
            continue

    return None


_input_profiles = os.getenv("SCRAPE_PROFILES", "").strip()

if _input_profiles:
    PROFILES_TO_SCRAPE = [
        item.strip()
        for item in _input_profiles.split(",")
        if item.strip()
    ]
    print("Run input: using the profile link(s) typed in the input box.")

_input_start_raw = os.getenv("SCRAPE_START_DATE", "")
_input_start = _date_from_input(_input_start_raw)

if _input_start_raw.strip() and _input_start is None:
    print("Run input: the start date could not be understood, "
          "so the date written in the code is being used.")
elif _input_start is not None:
    START_DATE = _input_start
    print("Run input: using the start date typed in the input box.")

_input_end_raw = os.getenv("SCRAPE_END_DATE", "")
_input_end = _date_from_input(_input_end_raw, end_of_day=True)

if _input_end_raw.strip() and _input_end is None:
    print("Run input: the end date could not be understood, "
          "so the date written in the code is being used.")
elif _input_end is not None:
    END_DATE = _input_end
    print("Run input: using the end date typed in the input box.")

# Credentials are NEVER written in the code. They are read from environment
# variables:
#   - On GitHub: from GitHub Secrets (TWITTER_USERNAME / TWITTER_PASSWORD),
#     passed in by the workflow file.
#   - On your own computer: from a local .env file (which is never uploaded)
#     or from environment variables you set yourself.
# The code simply reads the Secrets of the repository it is running in:
# a run inside the owner's repository logs in with the owner's ID/password;
# a run inside someone's own copy/fork logs in with that person's
# ID/password (they add the same two Secrets in their own repository).
try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

X_EMAIL = os.getenv("TWITTER_USERNAME", "")
X_PASSWORD = os.getenv("TWITTER_PASSWORD", "")

# True only when the script is running inside GitHub Actions.
IS_CI = os.getenv("GITHUB_ACTIONS") == "true"

# Automatic Chrome version detection (must stay None on GitHub: the runner
# has its own Chrome version, and a pinned number breaks the driver there).
CHROME_VERSION_MAIN = None

SCROLL_RATIO = 0.7
SCROLL_PAUSE = 1.3
MAX_IDLE_SCROLLS = 10
OLD_POSTS_TO_STOP = 5
LOGIN_TIMEOUT = 600
DEBUG = True

# ==============================================================================
# CREATE DIRECTORIES AND CONFIGURE TEMP FILES
# ==============================================================================

for folder in (
    BASE_DIR,
    OUTPUT_FOLDER,
    AUTOMATION_PROFILE_DIR,
    TEMP_DIR,
    DEBUG_FOLDER,
):
    os.makedirs(folder, exist_ok=True)

tempfile.tempdir = TEMP_DIR
os.environ["TMPDIR"] = TEMP_DIR
os.environ["TMP"] = TEMP_DIR
os.environ["TEMP"] = TEMP_DIR
os.environ["TMP"] = TEMP_DIR

import undetected_chromedriver as uc

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException, WebDriverException

STOP = False

# ==============================================================================
# HELPER FUNCTIONS
# ==============================================================================

def log(msg):
    line = f"[{datetime.now().strftime('%H:%M:%S')}] {msg}"
    try:
        print(line, flush=True)
    except UnicodeEncodeError:
        print(line.encode("ascii", "replace").decode(), flush=True)


NUM_LINE = re.compile(r"^\d[\d.,]*\s?[KkMmBb]?$")
HASH_RE = re.compile(r"(?<![\w/&#])#[\w\u0900-\u097F\u200c\u200d]+")
MENTION_RE = re.compile(r"(?<![\w.])@[A-Za-z0-9_]{1,15}")
ILLEGAL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def parse_count(text):
    """Convert 1.3K, 2.4M and 1,234 into integers."""

    if not text:
        return 0

    text = text.strip().replace(",", "").replace("\u00a0", "").replace(" ", "")
    match = re.search(r"([\d.]+)([KkMmBb]?)", text)

    if not match:
        return 0

    try:
        number = float(match.group(1))
    except ValueError:
        return 0

    multiplier = {
        "K": 1_000,
        "M": 1_000_000,
        "B": 1_000_000_000,
    }.get(match.group(2).upper(), 1)

    return int(round(number * multiplier))


def to_count(text):
    text = (text or "").strip()
    return parse_count(text) if NUM_LINE.match(text) else None


def count_from_label(label, word_regex):
    match = re.search(
        r"(\d[\d.,]*\s?[KkMm]?)\s+" + word_regex,
        label or "",
        re.I,
    )
    return parse_count(match.group(1)) if match else None


def metric(item, word_regex, btn_key):
    value = count_from_label(item.get("label"), word_regex)
    button = item.get(btn_key) or {}

    if value is None:
        value = count_from_label(button.get("label"), word_regex)

    if value is None:
        value = to_count(button.get("text"))

    return value if value is not None else 0


def strip_illegal(text):
    return ILLEGAL_RE.sub("", text or "")


def unique_ci(items):
    """Remove duplicates ignoring case, preserving original order."""

    seen = set()
    output = []

    for item in items:
        key = item.strip().lower()

        if key and key not in seen:
            seen.add(key)
            output.append(item.strip())

    return output


def parse_handle(value):
    """Accept username, @username, or full X/Twitter profile URL."""

    value = (value or "").strip()

    if not value:
        return ""

    if "x.com" in value or "twitter.com" in value:
        parsed = urlparse(
            value if "://" in value else "https://" + value
        )
        parts = [part for part in parsed.path.split("/") if part]
        return parts[0].lstrip("@") if parts else ""

    return value.lstrip("@").strip("/").split("/")[0]


def iso_to_local_datetime(iso):
    """Convert an ISO timestamp to the computer's local time."""

    if not iso:
        return None

    try:
        dt = datetime.fromisoformat(iso.replace("Z", "+00:00"))

        if dt.tzinfo is not None:
            dt = dt.astimezone()

        return dt.replace(tzinfo=None)

    except Exception:
        return None


def clean_filename(value):
    return re.sub(r"[^\w\-]+", "_", value).strip("_")[:40] or "profile"


# ==============================================================================
# EXCEL SAVING
# ==============================================================================

def save_to_excel(rows, page_name, out_path, quiet=False):

    if not rows:
        if not quiet:
            log("No posts were collected, so there is nothing to save.")
        return

    rows = sorted(rows, key=lambda row: row[0], reverse=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    thin = Side(style="thin")
    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )

    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    headers = [
        "Date",
        "Like",
        "Comment",
        "Repost",
        "Post Link",
        "Content",
        "Mentions",
        "Hashtags",
    ]

    ncols = len(headers)

    ws.merge_cells(
        start_row=1, start_column=1,
        end_row=1, end_column=ncols
    )
    ws["A1"] = page_name
    ws["A1"].font = Font(bold=True, size=11)

    ws.merge_cells(
        start_row=2, start_column=1,
        end_row=2, end_column=ncols
    )
    ws["A2"] = "Twitter"

    for col, header in enumerate(headers, start=1):
        ws.cell(row=3, column=col, value=header)

    def put_text(row, col, value):
        cell = ws.cell(row=row, column=col)
        cell.value = strip_illegal(value)[:32000]

        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"

        return cell

    for i, row in enumerate(rows, start=4):

        (
            dt,
            likes,
            comments,
            reposts,
            link,
            content,
            mentions,
            hashtags,
        ) = row

        ws.cell(row=i, column=1, value=dt.date()).number_format = "m/d/yyyy"
        ws.cell(row=i, column=2, value=likes)
        ws.cell(row=i, column=3, value=comments)
        ws.cell(row=i, column=4, value=reposts)

        link_cell = put_text(i, 5, link)

        if link:
            link_cell.hyperlink = link
            link_cell.font = Font(color="0563C1", underline="single")

        put_text(i, 6, content)
        put_text(i, 7, mentions)
        put_text(i, 8, hashtags)

    last_row = 3 + len(rows)

    for row_num in range(1, last_row + 1):
        for col_num in range(1, ncols + 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.border = border
            cell.alignment = (
                center if col_num <= 4 or row_num <= 3 else left
            )

    for col, width in zip(
        "ABCDEFGH",
        (14, 10, 10, 10, 48, 70, 30, 30),
    ):
        ws.column_dimensions[col].width = width

    try:
        wb.save(out_path)

    except PermissionError:
        alternative = out_path.replace(
            ".xlsx",
            f"_{datetime.now().strftime('%H%M%S')}.xlsx",
        )
        wb.save(alternative)
        log(f"Excel file was open; saved under: {alternative}")
        return

    if not quiet:
        log(f"Saved {len(rows)} rows -> {out_path}")


# ==============================================================================
# STOP HANDLER
# ==============================================================================

def signal_handler(sig, frame):
    global STOP
    log("Stop requested. Finishing the current step, then saving data...")
    STOP = True


# ==============================================================================
# CHROME DRIVER
# ==============================================================================

def _installed_chrome_major():
    """
    Return the major version of the Chrome/Chromium installed on this
    machine (for example 155), or None when it cannot be detected.
    Without this, undetected_chromedriver downloads the driver for the
    newest Chrome release, which breaks whenever the installed browser
    is one version behind (as on GitHub runners: driver 156 vs
    browser 155, and Chrome then refuses to start).
    """

    for command in (
        ["google-chrome", "--version"],
        ["google-chrome-stable", "--version"],
        ["chromium", "--version"],
        ["chromium-browser", "--version"],
    ):
        try:
            result = subprocess.run(
                command, capture_output=True, text=True, timeout=10
            )
        except Exception:
            continue

        match = re.search(r"(\d+)\.", result.stdout or "")

        if match:
            return int(match.group(1))

    return None


def setup_driver():

    options = uc.ChromeOptions()
    options.add_argument("--lang=en-US")

    if IS_CI:
        # GitHub runners have no screen: Chrome must run headless there,
        # and these flags are required on CI/container machines.
        options.add_argument("--headless=new")
        options.add_argument("--no-sandbox")
        options.add_argument("--disable-dev-shm-usage")
        options.add_argument("--disable-gpu")
        options.add_argument("--window-size=1920,1080")
    else:
        options.add_argument("--start-maximized")

    kwargs = {
        "options": options,
        "user_data_dir": AUTOMATION_PROFILE_DIR,
        "use_subprocess": True,
    }

    version_main = CHROME_VERSION_MAIN

    if version_main is None and IS_CI:
        # On GitHub, force the driver to match the browser that is
        # actually installed on the runner, not the newest release.
        version_main = _installed_chrome_major()

        if version_main is not None:
            log(f"Detected Chrome {version_main} on this machine; "
                "using the matching ChromeDriver.")

    if version_main is not None:
        kwargs["version_main"] = version_main

    return uc.Chrome(**kwargs)


def switch(driver, handle):

    driver.switch_to.window(handle)

    try:
        driver.execute_cdp_cmd("Page.bringToFront", {})
    except Exception:
        pass


# ==============================================================================
# LOGIN AUTOMATION
# ==============================================================================

def has_login_cookies(driver):

    try:
        names = {
            cookie.get("name")
            for cookie in driver.get_cookies()
        }
        return "auth_token" in names or "twid" in names

    except Exception:
        return False


def wait_for_manual_login(driver, timeout=LOGIN_TIMEOUT):
    """
    Wait for manual login or additional verification in Chrome.
    Does not attempt to bypass security checks.
    """

    if IS_CI:
        # On GitHub there is no screen and no person. Waiting 10 minutes
        # would only waste the run, so fail fast and say why.
        log("Manual login is not possible on GitHub (no screen, no person).")
        log("On GitHub the login has to succeed with ID/password alone.")
        return False

    log("Complete any required verification manually in Chrome.")

    deadline = time.time() + timeout

    while time.time() < deadline and not STOP:

        if has_login_cookies(driver):
            log("Login detected.")
            time.sleep(3)
            return True

        time.sleep(2)

    log("Login was not completed before the timeout.")
    return False


def _find_visible(driver, by, selector):
    """Return the first visible, enabled element matching the selector."""

    for element in driver.find_elements(by, selector):
        try:
            if element.is_displayed() and element.is_enabled():
                return element
        except Exception:
            continue
    return None


def _wait_visible(driver, by, selector, timeout=25):
    """Wait until a visible, enabled element appears, then return it."""

    deadline = time.time() + timeout

    while time.time() < deadline and not STOP:
        element = _find_visible(driver, by, selector)
        if element is not None:
            return element
        time.sleep(0.5)

    return None


def _click_button_with_text(driver, text, timeout=25):
    """Click the visible button whose label is exactly the given text."""

    deadline = time.time() + timeout

    while time.time() < deadline and not STOP:
        for button in driver.find_elements(By.XPATH, "//button"):
            try:
                if button.is_displayed() and \
                        (button.text or "").strip().lower() == text.lower():
                    button.click()
                    return True
            except Exception:
                continue
        time.sleep(0.5)

    return False


def _save_login_debug(driver, tag):
    """
    Save a photo and the HTML of the page the driver is currently on
    into the Debug folder, so a failed login on GitHub can be checked
    afterwards (the workflow uploads the Debug folder as an artifact).
    Also logs the page title and address.
    """

    try:
        stamp = datetime.now().strftime("%H%M%S")
        photo = os.path.join(DEBUG_FOLDER, f"login_{tag}_{stamp}.png")
        driver.save_screenshot(photo)
        log(f"Saved a photo of the page for checking: {photo}")

        page_file = os.path.join(DEBUG_FOLDER, f"login_{tag}_{stamp}.html")
        with open(page_file, "w", encoding="utf-8") as handle:
            handle.write(driver.page_source or "")
        log(f"Saved the page details for checking: {page_file}")
    except Exception:
        pass

    try:
        log(f"Page title: {driver.title!r}")
        log(f"Page address: {driver.current_url}")
    except Exception:
        pass


def wait_for_login(driver):
    """
    Log in with ID/password only, so anyone can run this code with their
    own account:
    1. If the saved Chrome profile is already logged in, use it.
    2. Otherwise fill the real X login form: enter the login ID (username
       or email) on https://x.com/i/flow/login, click Next, handle the
       extra confirmation step X sometimes shows by entering the login
       ID once more, then enter the password and click Log in.
    3. If X demands a step only a human can do (for example a code sent
       to email/phone), fall back to manual completion - possible only
       on your own computer, never on GitHub.
    """

    if has_login_cookies(driver):
        log("Already logged in using the saved Chrome profile.")
        return True

    if not X_EMAIL or not X_PASSWORD:
        log("TWITTER_USERNAME or TWITTER_PASSWORD is missing "
            "in THIS repository.")
        log("Add both under: Settings > Secrets and variables > Actions, "
            "then run again.")
        log("(In the owner's repo these are the owner's login; in your own "
            "copy/fork, add YOUR login there.)")
        return wait_for_manual_login(driver)

    try:
        # STEP 1: Open the login page and find the login ID field
        # (username / email / phone). On a cloud machine the page can
        # load slowly or come up blank, so wait longer and load the
        # page a second time before giving up.
        id_input = None

        for attempt in (1, 2):
            driver.get("https://x.com/i/flow/login")
            time.sleep(2)
            id_input = _wait_visible(
                driver, By.CSS_SELECTOR, 'input[name="text"]', timeout=45
            )
            if id_input is not None:
                break
            log("The login form did not appear; trying once more...")

        if id_input is None:
            _save_login_debug(driver, "no_form")
            raise TimeoutException("The login ID field was not found.")

        id_input.click()
        id_input.clear()
        id_input.send_keys(X_EMAIL)
        log("Login ID entered.")

        # STEP 2: Click Next.
        if not _click_button_with_text(driver, "Next"):
            raise TimeoutException("The 'Next' button was not found.")

        # STEP 3: Wait for the password field. X sometimes first shows an
        # extra confirmation screen (enter username/phone again) - handle
        # it by entering the same login ID once more.
        password_input = None
        deadline = time.time() + 30

        while time.time() < deadline and not STOP:
            password_input = _find_visible(
                driver, By.CSS_SELECTOR, 'input[name="password"]'
            )
            if password_input is not None:
                break

            extra_input = _find_visible(
                driver, By.CSS_SELECTOR, 'input[name="text"]'
            )
            if extra_input is not None:
                log("X asked for one extra confirmation; "
                    "entering the login ID again.")
                extra_input.click()
                extra_input.clear()
                extra_input.send_keys(X_EMAIL)
                _click_button_with_text(driver, "Next", timeout=5)

            time.sleep(0.5)

        if password_input is None:
            raise TimeoutException("The password field was not found.")

        # STEP 4: Fill the password field.
        password_input.click()
        password_input.clear()
        password_input.send_keys(X_PASSWORD)
        log("Password entered.")

        # STEP 5: Click Log in.
        if not _click_button_with_text(driver, "Log in"):
            raise TimeoutException("The 'Log in' button was not found.")

        log("Submitted ID/password. Waiting for login confirmation...")

        deadline = time.time() + LOGIN_TIMEOUT

        while time.time() < deadline and not STOP:

            if has_login_cookies(driver):
                log("Login detected.")
                time.sleep(3)
                return True

            time.sleep(2)

        log("Login was not confirmed automatically.")
        _save_login_debug(driver, "not_confirmed")
        return wait_for_manual_login(driver)

    except (TimeoutException, WebDriverException) as exc:

        log(
            "Automated login could not complete: "
            f"{exc.__class__.__name__}."
        )
        log(
            "The login page may have changed or may require an additional step."
        )

        return wait_for_manual_login(driver)


# ==============================================================================
# JAVASCRIPT TO READ POSTS FROM THE VISIBLE X PAGE
# ==============================================================================

HAS_TWEET_JS = r"""
return !!document.querySelector('article[data-testid="tweet"]');
"""

TWEET_HELPERS_JS = r"""
function cleanText(el) {
  var c = el.cloneNode(true);

  c.querySelectorAll('img[alt]').forEach(function(im) {
    im.replaceWith(document.createTextNode(im.getAttribute('alt') || ''));
  });

  c.querySelectorAll('br').forEach(function(b) {
    b.replaceWith(document.createTextNode('\n'));
  });

  return (c.textContent || '').trim();
}

function btnInfo(art, ids) {
  for (var i = 0; i < ids.length; i++) {
    var b = art.querySelector('[data-testid="' + ids[i] + '"]');

    if (b) {
      return {
        label: b.getAttribute('aria-label') || '',
        text: (b.innerText || '').trim()
      };
    }
  }

  return {label: '', text: ''};
}

function readTweet(art, handle) {
  var timeEl = null;
  var link = null;
  var times = art.querySelectorAll('time[datetime]');

  for (var i = 0; i < times.length; i++) {
    var a = times[i].closest('a[href*="/status/"]');

    if (a) {
      timeEl = times[i];
      link = a;
      break;
    }
  }

  if (!timeEl) return null;

  var m = (link.getAttribute('href') || '')
      .match(/(?:^|\/)([^\/?#]+)\/status\/(\d+)/);

  if (!m) return null;

  var ctx = '';

  art.querySelectorAll('[data-testid="socialContext"]').forEach(function(e) {
    ctx += ' ' + (e.textContent || '');
  });

  var tEls = art.querySelectorAll('[data-testid="tweetText"]');
  var tEl = null;

  for (var j = 0; j < tEls.length && !tEl; j++) {
    var p = tEls[j].parentElement;
    var inLink = false;

    while (p && p !== art) {
      if (p.getAttribute('role') === 'link') {
        inLink = true;
        break;
      }
      p = p.parentElement;
    }

    if (!inLink) tEl = tEls[j];
  }

  var text = '';
  var tags = [];
  var ments = [];

  if (tEl) {
    text = cleanText(tEl);

    tEl.querySelectorAll('a[href]').forEach(function(a) {
      var h = a.getAttribute('href') || '';
      var tx = (a.textContent || '').replace(/\s+/g, ' ').trim();

      if (!tx) return;

      if (tx.charAt(0) === '#' || /\/hashtag\//.test(h)) {
        tags.push(tx.charAt(0) === '#' ? tx : '#' + tx);
      } else if (tx.charAt(0) === '@') {
        ments.push(tx);
      }
    });
  }

  var label = '';
  var groups = art.querySelectorAll('[role="group"][aria-label]');

  for (var k = 0; k < groups.length; k++) {
    var lb = groups[k].getAttribute('aria-label') || '';

    if (/repl|repost|like/i.test(lb)) {
      label = lb;
      break;
    }
  }

  return {
    id: m[2],
    author: m[1],
    iso: timeEl.getAttribute('datetime') || '',
    pinned: /pinned/i.test(ctx),
    repost: /repost|retweet/i.test(ctx),
    other: m[1].toLowerCase() !== handle,
    text: text,
    truncated: !!art.querySelector('[data-testid="tweet-text-show-more-link"]'),
    tags: tags,
    mentions: ments,
    label: label,
    replyBtn: btnInfo(art, ['reply']),
    repostBtn: btnInfo(art, ['retweet', 'unretweet']),
    likeBtn: btnInfo(art, ['like', 'unlike'])
  };
}
"""

COLLECT_TWEETS_JS = TWEET_HELPERS_JS + r"""
var handle = (arguments[0] || '').toLowerCase();
var out = [];

document.querySelectorAll('article[data-testid="tweet"]').forEach(function(art) {
  try {
    var r = readTweet(art, handle);
    if (r) out.push(r);
  } catch (e) {}
});

return out;
"""

READ_POST_JS = TWEET_HELPERS_JS + r"""
var handle = (arguments[0] || '').toLowerCase();
var id = arguments[1] || '';
var arts = document.querySelectorAll('article[data-testid="tweet"]');

for (var i = 0; i < arts.length; i++) {
  try {
    var r = readTweet(arts[i], handle);
    if (r && r.id === id) return r;
  } catch (e) {}
}

return null;
"""

CLICK_RETRY_JS = r"""
var c = 0;

document.querySelectorAll('[role="button"], button').forEach(function(b) {
  var t = (b.innerText || '').trim().toLowerCase();

  if (t === 'retry') {
    try {
      b.click();
      c++;
    } catch (e) {}
  }
});

return c;
"""


# ==============================================================================
# SCRAPING FUNCTIONS
# ==============================================================================

def fetch_full_post(driver, grid, work, handle, tweet_id):
    """Open a long post in another tab to retrieve its full text."""

    result = None
    switch(driver, work)

    try:
        driver.get(f"https://x.com/{handle}/status/{tweet_id}")
        deadline = time.time() + 20

        while time.time() < deadline:

            try:
                data = driver.execute_script(
                    READ_POST_JS, handle, tweet_id
                )
            except Exception:
                data = None

            if data and data.get("text"):
                result = data
                break

            time.sleep(1)

    except Exception:
        result = None

    finally:
        switch(driver, grid)

    return result


def get_profile_name(driver, handle):

    title = driver.title or ""
    match = re.match(
        r"^\s*(?:\(\d+\+?\)\s*)?(.*?)\s*\(@([^)]+)\)",
        title,
    )

    name = match.group(1).strip() if match else ""

    return f"{name} (@{handle})" if name else f"@{handle}"


def open_profile(driver, handle):
    """Open profile and wait until the first post is visible."""

    url = f"https://x.com/{handle}"

    for attempt in range(2):
        driver.get(url)

        try:
            WebDriverWait(driver, 30).until(
                lambda d: d.execute_script(HAS_TWEET_JS)
            )
            return True

        except Exception:
            if attempt == 0:
                log("Timeline did not load. Retrying...")
                driver.get("https://x.com/home")
                time.sleep(3)

    return False


def scrape_profile(driver, handle, page_name, out_path):

    rows = []
    known_ids = set()
    old_ids = set()
    idle = 0
    debug_done = False

    grid = driver.current_window_handle
    driver.switch_to.new_window("tab")
    work = driver.current_window_handle
    switch(driver, grid)

    try:
        while not STOP:

            try:
                items = driver.execute_script(
                    COLLECT_TWEETS_JS, handle
                ) or []
            except Exception:
                items = []

            new_found = False

            for item in items:

                if STOP:
                    break

                tweet_id = item["id"]

                if tweet_id in known_ids:
                    continue

                known_ids.add(tweet_id)
                new_found = True

                if DEBUG and not debug_done:
                    debug_done = True
                    shown = dict(item)
                    shown["text"] = (shown.get("text") or "")[:80]
                    log(f"DEBUG raw data of first post: {shown}")

                # Skip reposts and posts authored by other accounts.
                if item.get("repost") or item.get("other"):
                    continue

                dt = iso_to_local_datetime(item.get("iso"))

                if dt is None:
                    continue

                if dt > END_DATE:
                    continue

                if dt < START_DATE:
                    if not item.get("pinned"):
                        old_ids.add(tweet_id)
                    continue

                content = item.get("text") or ""
                tags = list(item.get("tags") or [])
                mentions = list(item.get("mentions") or [])

                if item.get("truncated"):

                    full = fetch_full_post(
                        driver, grid, work, handle, tweet_id
                    )

                    if full and full.get("text"):
                        content = full["text"]
                        tags = list(full.get("tags") or [])
                        mentions = list(full.get("mentions") or [])

                content = strip_illegal(content).strip()

                tags += [
                    match.group(0)
                    for match in HASH_RE.finditer(content)
                ]

                tags = unique_ci([
                    strip_illegal(tag)
                    for tag in tags
                    if not re.fullmatch(r"#\d+", tag)
                ])

                mentions += [
                    match.group(0)
                    for match in MENTION_RE.finditer(content)
                ]

                mentions = unique_ci([
                    strip_illegal(mention)
                    for mention in mentions
                ])

                if not content:
                    content = "Media Only / No Text"

                likes = metric(item, r"likes?", "likeBtn")
                comments = metric(item, r"repl(?:y|ies)", "replyBtn")
                reposts = metric(item, r"reposts?", "repostBtn")

                link = (
                    f"https://x.com/{item['author']}/status/{tweet_id}"
                )

                rows.append((
                    dt,
                    likes,
                    comments,
                    reposts,
                    link,
                    content,
                    ", ".join(mentions),
                    ", ".join(tags),
                ))

                preview = content[:40].replace("\n", " ")

                log(
                    f"Post #{len(rows):<4} "
                    f"{dt.strftime('%m/%d/%Y %H:%M')} | "
                    f"Likes: {likes} | Comments: {comments} | "
                    f"Reposts: {reposts} | /status/{tweet_id} | {preview}"
                )

                # Save every 10 posts.
                if len(rows) % 10 == 0:
                    save_to_excel(
                        rows, page_name, out_path, quiet=True
                    )

            if STOP:
                break

            if len(old_ids) >= OLD_POSTS_TO_STOP:
                log("Reached posts older than START_DATE. Stopping profile.")
                break

            if new_found:
                idle = 0
            else:
                idle += 1
                log(f"No new posts in view ({idle}/{MAX_IDLE_SCROLLS}).")

                if idle >= MAX_IDLE_SCROLLS:
                    log("Reached the end of the profile.")
                    break

            try:
                info = driver.execute_script(
                    "var y=window.scrollY;"
                    "window.scrollBy(0, Math.floor(window.innerHeight*%s));"
                    "return [y, window.scrollY];" % SCROLL_RATIO
                )
                at_bottom = bool(info) and info[0] == info[1]

            except Exception:
                at_bottom = False

            time.sleep(SCROLL_PAUSE)

            if idle >= 1:
                try:
                    if driver.execute_script(CLICK_RETRY_JS):
                        log("Clicked 'Retry' on the timeline.")
                        time.sleep(2)
                except Exception:
                    pass

            if at_bottom and idle >= 2:
                time.sleep(2)

                try:
                    driver.execute_script("window.scrollBy(0, -700);")
                    time.sleep(1.0)

                    driver.execute_script(
                        "window.scrollTo(0, document.documentElement.scrollHeight);"
                    )

                    time.sleep(SCROLL_PAUSE)

                except Exception:
                    pass

    finally:
        # Save everything collected when the profile ends or is interrupted.
        save_to_excel(rows, page_name, out_path)

        try:
            driver.switch_to.window(work)
            driver.close()
        except Exception:
            pass

        try:
            driver.switch_to.window(grid)
        except Exception:
            pass

    return len(rows)


# ==============================================================================
# MAIN
# ==============================================================================

def main():

    global STOP
    STOP = False

    handles = [
        parse_handle(profile)
        for profile in PROFILES_TO_SCRAPE
    ]
    handles = [handle for handle in handles if handle]

    if not handles:
        log("No usernames found in PROFILES_TO_SCRAPE.")
        return

    if START_DATE > END_DATE:
        log("START_DATE is later than END_DATE. Please fix the dates.")
        return

    log(f"Base directory: {BASE_DIR}")
    log(f"Excel output: {OUTPUT_FOLDER}")
    log(f"Chrome profile: {AUTOMATION_PROFILE_DIR}")
    log(f"Temporary files: {TEMP_DIR}")
    log(f"Debug screenshots: {DEBUG_FOLDER}")
    log(
        f"Date range: {START_DATE:%Y-%m-%d %H:%M} "
        f"to {END_DATE:%Y-%m-%d %H:%M}"
    )
    log(f"Profiles in this run: {', '.join(PROFILES_TO_SCRAPE)}")

    try:
        signal.signal(signal.SIGINT, signal_handler)
    except Exception:
        pass

    driver = None

    try:
        driver = setup_driver()
        driver.get("https://x.com/home")
        time.sleep(4)

        if not wait_for_login(driver):
            log("Login did not complete. Please run the script again.")
            if IS_CI:
                # On GitHub, a failed login must fail the run (red cross),
                # not silently look successful.
                raise SystemExit(1)
            return

        print("\n" + "=" * 62)
        print(" Scraping started. Press Ctrl+C to stop and save.")
        print(" Do not close the browser window while it is running.")
        print("=" * 62 + "\n")

        for handle in handles:

            if STOP:
                break

            log(f"Opening profile: @{handle}")

            if not open_profile(driver, handle):
                log(f"No posts could be loaded for @{handle}. Skipping.")

                try:
                    screenshot = os.path.join(
                        DEBUG_FOLDER,
                        f"debug_x_{clean_filename(handle)}.png",
                    )
                    driver.save_screenshot(screenshot)
                    log(f"Saved debug screenshot: {screenshot}")

                except Exception:
                    pass

                continue

            time.sleep(2)

            page_name = get_profile_name(driver, handle)
            log(f"Profile: {page_name}")

            filename = (
                f"twitter_{clean_filename(handle)}_"
                f"{START_DATE:%Y%m%d}_to_{END_DATE:%Y%m%d}_"
                f"{datetime.now():%H%M%S}.xlsx"
            )

            out_path = os.path.join(OUTPUT_FOLDER, filename)

            count = scrape_profile(
                driver, handle, page_name, out_path
            )

            if count:
                log(f"Done with @{handle}: {count} posts.")
                log(f"Excel file: {out_path}")
            else:
                log(f"No posts found for @{handle} in this date range.")

        log(
            "ALL PROFILES FINISHED!"
            if not STOP
            else "Stopped. Collected data so far has been saved."
        )

    except KeyboardInterrupt:
        log("Interrupted. Collected data so far has been saved.")

    finally:
        if driver is not None:
            try:
                driver.quit()
            except Exception:
                pass


if __name__ == "__main__":
    main()
